"""Real local multi-agent smoke test; uses the app's .env through the server.

Run: .venv/bin/python smoke_multiagent.py --document-type proposal
No keys or source documents are printed. The synthetic case is checked and
exported to DOCX, PDF and ZIP so the local pipeline is verified end to end.
"""
import argparse
import io
import json
import zipfile
import time
from pathlib import Path

import requests
from docx import Document


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", default="http://127.0.0.1:8001")
    parser.add_argument("--document-type", choices=["proposal", "brd"], default="proposal")
    parser.add_argument("--out", default="out/multiagent-smoke")
    parser.add_argument("--thread-id", help="Continue/check an existing synthetic case")
    args = parser.parse_args()
    base = args.base_url.rstrip("/")
    session = requests.Session()
    fixture = Path(__file__).parent / "tests" / "fixtures" / "multiagent_requirement.txt"
    if args.thread_id:
        response = session.get(base + "/api/agent/state/" + args.thread_id, timeout=30)
    else:
        with fixture.open("rb") as source:
            response = session.post(base + "/api/agent/start", data={"document_type": args.document_type},
                                    files={"files": (fixture.name, source, "text/plain")}, timeout=300)
    response.raise_for_status(); first = response.json()
    thread = first["thread_id"]
    if first.get("stage") == "clarification":
        answers = {q["id"]: "Leave open for discovery" for q in first["payload"]["questions"]}
        response = session.post(base + "/api/agent/resume", json={
            "thread_id": thread, "response": {"answers": answers}
        }, timeout=900)
    elif first.get("retryable"):
        response = session.post(base + "/api/agent/retry", json={"thread_id": thread}, timeout=900)
    for attempt in range(3):
        if response.status_code != 503:
            break
        print("Model rate limit encountered; continuing saved nodes after a short wait.")
        time.sleep(25)
        response = session.post(base + "/api/agent/retry", json={"thread_id": thread}, timeout=900)
    if not response.ok:
        raise RuntimeError(f"Specialist generation failed ({response.status_code}): {response.text[:500]}")
    second = response.json()
    if second.get("stage") == "draft_review":
        review = second["payload"]
        assert review["quality_report"]["passed"]
        assert set(review["specialist_findings"]) == {"solution", "delivery", "commercial", "risk"}
        response = session.post(base + "/api/agent/resume", json={
            "thread_id": thread, "response": {"decision": "accept", "reviewer": "Smoke test reviewer"}
        }, timeout=300)
    response.raise_for_status(); final = response.json()
    assert final["status"] == "reviewed"
    output = Path(args.out); output.mkdir(parents=True, exist_ok=True)
    for format in ["docx", "pdf", "zip"]:
        exported = session.get(base + f"/api/agent/export/{thread}?format={format}", timeout=300)
        exported.raise_for_status()
        name = Path(exported.headers["x-proposal-filename"]).name
        (output / name).write_bytes(exported.content)
        if format == "docx":
            doc = Document(io.BytesIO(exported.content))
            text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
                cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            assert "Solution Recommendations" in text and "Requirement Traceability" in text
        elif format == "pdf":
            assert exported.content.startswith(b"%PDF")
        else:
            with zipfile.ZipFile(io.BytesIO(exported.content)) as archive:
                assert len(archive.namelist()) == 3
                audit = json.loads(archive.read(next(name for name in archive.namelist() if name.endswith(".json"))))
                assert audit["agent_audit"]["quality_report"]["passed"]
        print(f"{format.upper()} export verified: {name}")
    print(f"Multi-agent {args.document_type} passed: {len(final['trace'])} trace steps; thread {thread}")
    print(f"Reviewable sample outputs: {output.resolve()}")


if __name__ == "__main__":
    main()
