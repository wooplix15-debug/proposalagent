import copy
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from langgraph.checkpoint.memory import InMemorySaver

import multi_agent_graph as graph
import multi_agent_specialists as specialists
import wooplix_agent as brand
from multi_agent_contracts import SolutionReport
from multi_agent_tools import apply_text_patches, attach_findings, check_references, critic_batches, requirement_register


def context():
    return {
        "source": "Client brief.txt", "requirement": "Zoho CRM\nLead management",
        "analysis": {"summary": "CRM required", "requirement_sections": [
            {"product": "Zoho CRM", "requirements": ["Lead management"]}
        ], "questions": [], "duration_estimates": [], "commercial_categories": []},
        "answers": {}, "completed": [], "timing_sources": [],
    }


def proposal():
    return {"client": {"company_name": "Example Client", "project_name": "CRM", "contact": ""},
            "project_introduction": "Implement Zoho CRM.",
            "scope": [{"product": "Zoho CRM", "areas": [{"area": "Leads", "tasks": ["Lead management"]}]}],
            "prerequisites": [], "deliverables": ["Configured CRM."], "open_points": [],
            "status": "DRAFT"}


class MultiAgentGroundingTests(unittest.TestCase):
    def test_large_brd_critic_packets_have_bounded_input_and_original_paths(self):
        document = {"summary": "A long proposed explanation. " * 700,
                    "specification_tables": [
                        {"origin": "source", "rows": [["Source data must not be rewritten"]]},
                        {"origin": "proposed", "rows": [["Suggested lead assignment process"]]},
                    ]}
        packets = critic_batches(document)
        self.assertGreater(len(packets), 1)
        for packet in packets:
            self.assertLessEqual(sum(len(field["path"]) + len(field["text"]) + 32 for field in packet), 7000)
        paths = {field["path"] for packet in packets for field in packet}
        self.assertIn("specification_tables[1].rows[0][0]", paths)
        self.assertNotIn("specification_tables[0].rows[0][0]", paths)

    def test_register_ids_remain_stable_when_checklist_reordered(self):
        ctx = context()
        ctx["analysis"]["requirement_sections"][0]["requirements"].append("Follow-up reminders")
        first = requirement_register(ctx)
        ctx["analysis"]["requirement_sections"][0]["requirements"].reverse()
        second = requirement_register(ctx)
        self.assertEqual({r["requirement"]: r["id"] for r in first},
                         {r["requirement"]: r["id"] for r in second})
        self.assertEqual(first[0]["source"], "Client brief.txt")

    def test_unrequested_products_and_fabricated_evidence_are_rejected(self):
        ctx = context(); ctx["requirement_register"] = requirement_register(ctx)
        valid = {"solutions": [{"requirement_ids": [ctx["requirement_register"][0]["id"]],
                                "product": "Zoho CRM", "evidence_ids": []}]}
        check_references(valid, ctx)
        bad = copy.deepcopy(valid); bad["solutions"][0]["product"] = "Zoho Books"
        with self.assertRaises(ValueError):
            check_references(bad, ctx)
        bad = copy.deepcopy(valid); bad["solutions"][0]["design"] = "Guarantee a 99.9 percent SLA"
        with self.assertRaises(ValueError):
            check_references(bad, ctx)
        bad = copy.deepcopy(valid); bad["solutions"][0]["evidence_ids"] = ["invented"]
        with self.assertRaises(ValueError):
            check_references(bad, ctx)

    def test_repair_cannot_change_pricing_source_rows_or_requirement_ids(self):
        doc = {"commercials": {"total": ""}, "requirements": [{"id": "BR-001"}],
               "specification_tables": [{"origin": "source", "rows": [["Phone", "Text"]]}]}
        for path in ("commercials.total", "requirements[0].id", "specification_tables[0].rows[0][1]"):
            with self.assertRaises(ValueError):
                apply_text_patches(doc, [{"path": path, "replacement": "Changed"}])
        self.assertEqual(apply_text_patches({"deliverables": ["Guaranteed result"]},
                         [{"path": "deliverables[0]", "replacement": "Result subject to review"}])["deliverables"],
                         ["Result subject to review"])

    def test_malformed_specialist_output_is_not_reported_as_success(self):
        ctx = context(); ctx["requirement_register"] = requirement_register(ctx)
        response = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='{"solutions":"wrong"}'))])
        with patch.object(specialists.workflow, "_groq_call_with_fallback", return_value=response) as call:
            with self.assertRaises(specialists.workflow.AnalysisServiceError):
                specialists._json_agent("test role", "Return findings", {}, SolutionReport, ctx)
            self.assertEqual(call.call_count, 2)

    def test_unverifiable_critic_excerpt_is_retried_before_acceptance(self):
        ctx = context(); ctx["requirement_register"] = requirement_register(ctx)
        issue = {"severity": "blocking", "category": "unsupported_claim", "path": "deliverables[0]",
                 "excerpt": "Text that is not in the document", "reason": "Incorrect claim",
                 "suggested_fix": "Qualify it", "requirement_ids": [ctx["requirement_register"][0]["id"]]}
        import json
        def response(value):
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(value)))])
        with patch.object(specialists.workflow, "_groq_call_with_fallback",
                          side_effect=[response({"issues": [issue]}), response({"issues": []})]) as call:
            result = specialists.critic_specialist(ctx, "proposal", proposal(), [])
            self.assertTrue(result["passed"])
            self.assertEqual(call.call_count, 2)

    def test_planning_findings_are_in_word_and_pdf_html(self):
        ctx = context(); register = requirement_register(ctx)
        findings = {
            "solution": {"solutions": [{"product": "Zoho CRM", "design": "Review lead ownership rules.",
                          "reason": "Lead management requested", "status": "proposed",
                          "requirement_ids": [register[0]["id"]]}], "open_decisions": []},
            "delivery": {"dependencies": [], "open_decisions": []},
            "commercial": {"dependencies": []},
            "risk": {"risks": [{"description": "Ownership rules are not confirmed.", "impact": "Assignment design is open.",
                                 "mitigation": "Confirm with the sales team.", "status": "open"}]},
        }
        doc = attach_findings(proposal(), findings, register)
        html = brand.build_html(doc)
        self.assertIn("Solution Recommendations", html)
        self.assertIn("Requirement Traceability", html)
        with tempfile.TemporaryDirectory() as work:
            path = Path(work) / "review.docx"; brand.build_docx(doc, str(path))
            output = Document(path)
            text = "\n".join(p.text for p in output.paragraphs) + "\n".join(
                cell.text for table in output.tables for row in table.rows for cell in row.cells)
            self.assertIn("Ownership rules are not confirmed.", text)
            self.assertIn(register[0]["id"], text)

    def test_stateless_graph_cannot_loop_past_one_critic_revision(self):
        ctx = context()
        plan = {"tasks": [{"specialist": name, "focus": "Review", "requirement_ids": []}
                          for name in ("solution", "delivery", "commercial", "risk")], "quality_focus": []}
        issue = {"severity": "blocking", "reason": "Unsupported guarantee", "path": "deliverables[0]",
                 "excerpt": "Configured CRM.", "category": "unsupported_claim",
                 "suggested_fix": "Qualify it", "requirement_ids": []}
        with patch.object(graph, "supervisor_specialist", return_value=plan), \
             patch.object(graph, "solution_specialist", return_value={"solutions": [], "open_decisions": []}), \
             patch.object(graph, "delivery_specialist", return_value={"timeline": {}, "dependencies": [], "open_decisions": []}), \
             patch.object(graph, "commercial_specialist", return_value={"commercials": None, "dependencies": []}), \
             patch.object(graph, "risk_specialist", return_value={"risks": []}), \
             patch.object(graph.brand, "draft_proposal", return_value=proposal()), \
             patch.object(graph, "critic_specialist", return_value={"passed": False, "issues": [issue]}) as critic, \
             patch.object(graph, "repair_specialist", return_value={"patches": [
                 {"path": "deliverables[0]", "replacement": "Configured CRM for review."}
             ]}) as repair:
            with self.assertRaises(graph.base_graph.AgentGraphError):
                graph.run_proposal(ctx, {})
            self.assertEqual(repair.call_count, 1)
            self.assertEqual(critic.call_count, 2)
