"""Independent agent for creating a source-grounded Business Requirements Document."""
from __future__ import annotations

import html
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import wooplix_agent as brand
from brd_structure import SPECIFICATION_SECTIONS, source_tables, statement_tables, proposed_design_tables


def _section_items(sections):
    """Assign stable BRD IDs to every distinct source requirement."""
    rows = []
    seen = set()
    number = 1
    for section in sections or []:
        product = str(section.get("product") or "Business requirement").strip()
        for requirement in section.get("requirements") or []:
            text = re.sub(r"\s+", " ", str(requirement)).strip()
            key = (product.casefold(), text.casefold())
            if not text or key in seen:
                continue
            seen.add(key)
            rows.append({"id": f"BR-{number:03}", "area": product, "requirement": text})
            number += 1
    return rows



def _source_configuration_table(area, requirements):
    """Keep a BRD usable if the model omits a small or secondary work area."""
    rows = [[row["requirement"], ""]
            for row in requirements[:60] if row.get("requirement")]
    return {
        "section": "Business Process",
        "title": f"{area} configuration requirements",
        "columns": ["Requirement", "Configuration detail"],
        "rows": rows,
        "requirement_ids": [row["id"] for row in requirements],
        "origin": "source",
    }


def _is_parallel_schedule_question(question):
    text = str(question or "").casefold()
    return (any(term in text for term in ("one after another", "at the same time", "overlap", "sequenc"))
            or ("schedule" in text and any(term in text for term in
                                           ("work areas", "work area", "work to be", "all work"))))


def _source_list(requirement, section, next_headers):
    """Read literal list items beneath one source heading."""
    lines = requirement.splitlines()
    start = next((i for i, line in enumerate(lines)
                  if re.match(rf"^\s*(?:\d+(?:\.\d+)*[.)]?\s*)?{section}\b", line, re.I)), None)
    if start is None:
        return []
    result = []
    for line in lines[start + 1:]:
        numbered_heading = re.match(r"^\s*\d+(?:\.\d+)*[.)]?\s+[A-Z]", line)
        known_heading = any(re.match(rf"^\s*(?:\d+[.)]\s*)?{header}\b", line, re.I)
                            for header in next_headers)
        if numbered_heading or known_heading:
            break
        text = re.sub(r"^\s*(?:[-•*]|\d+[.)])\s*", "", line).strip()
        if text:
            result.append(text)
    return list(dict.fromkeys(result))


def _acceptance_table_rows(requirement):
    """Read acceptance criteria from flattened DOCX/PDF tables when present."""
    lines = requirement.splitlines()
    header = next((i for i, line in enumerate(lines)
                   if "acceptance criterion" in line.casefold()
                   and "verified by" in line.casefold()), None)
    if header is None:
        return []
    criteria = []
    for line in lines[header + 1:]:
        low = line.casefold()
        if any(marker in low for marker in ("name & designation", "name and designation")):
            break
        cells = [cell.strip() for cell in line.split("|")]
        if len(cells) < 3 or not cells[0].isdigit():
            continue
        criterion = re.sub(r"\s+", " ", cells[1]).strip()
        if criterion:
            criteria.append(criterion)
    return list(dict.fromkeys(criteria))


def _combine_explicit_shared_areas(requirement, sections):
    """Merge one shared source checklist when it explicitly names two products."""
    raw = requirement.casefold()
    pairs = [
        ("Zoho Marketing Automation", "Zoho Campaigns",
         bool(re.search(r"zoho\s+marketing\s+automation\s*[,/&]+\s*(?:zoho\s+)?campaign", raw))),
        ("Zoho Creator", "Zoho Backstage",
         bool(re.search(r"zoho\s+creator\s*[,/&]+\s*(?:zoho\s+)?backstage", raw))),
    ]
    combined, consumed = [], set()
    for section in sections:
        product = str(section.get("product") or "").strip()
        if product in consumed:
            continue
        shared_pair = next((pair for pair in pairs if pair[2] and product in pair[:2]), None)
        if shared_pair:
            names = shared_pair[:2]
            peers = [row for row in sections if row.get("product") in names]
            values = [list(dict.fromkeys(str(item).strip() for item in peer.get("requirements", [])
                                        if str(item).strip())) for peer in peers]
            if len(peers) == 2 and values[0] == values[1]:
                combined.append({"product": " & ".join(names), "requirements": values[0]})
                consumed.update(names)
                continue
        combined.append(section)
    return combined


def _groq_json(system, payload, max_tokens=6500):
    """Run one structured BRD analysis pass, trying configured keys and models."""
    from groq import Groq, RateLimitError

    keys = brand.GROQ_API_KEYS or ([brand.GROQ_API_KEY] if brand.GROQ_API_KEY else [])
    if not keys:
        raise RuntimeError("The BRD agent needs a configured Groq API key.")
    models = list(dict.fromkeys([brand.GROQ_MODEL, "openai/gpt-oss-120b",
                                 "openai/gpt-oss-20b", "qwen/qwen3.8-27b"]))
    last_error = None
    for model in models:
        for key in keys:
            try:
                kwargs = {"model": model,
                          "messages": [{"role": "system", "content": system},
                                       {"role": "user", "content": json.dumps(payload, ensure_ascii=False)}],
                          "temperature": 0.1, "max_completion_tokens": max_tokens,
                          "response_format": {"type": "json_object"}}
                if "gpt-oss" in model:
                    kwargs["reasoning_effort"] = "low"
                response = Groq(api_key=key, timeout=55, max_retries=0).chat.completions.create(**kwargs)
                data = json.loads(response.choices[0].message.content or "{}")
                if isinstance(data, list):
                    data = {"specification_tables": data}
                if not isinstance(data, dict):
                    raise ValueError("The model did not return a JSON object")
                return data
            except RateLimitError as exc:
                last_error = exc
                match = re.search(r"try again in ([\d.]+)s", str(exc), re.I)
                if match:
                    time.sleep(min(float(match.group(1)) + 1, 20))
                break
            except Exception as exc:
                # Groq may report an oversized prompt as HTTP 413 rather than
                # RateLimitError. Do not mask that cause as a generic failure.
                if getattr(exc, "status_code", None) == 413:
                    raise RuntimeError(
                        "This requirement is too detailed for one BRD analysis request. "
                        "Please split it into two documents and try again."
                    ) from exc
                last_error = exc
    if isinstance(last_error, RateLimitError):
        match = re.search(r"try again in ([\d.]+)s", str(last_error), re.I)
        wait = f" Please wait about {max(1, round(float(match.group(1))))} seconds and try again." if match else " Please try again in a moment."
        raise RuntimeError("The BRD analysis service is temporarily busy." + wait) from last_error
    raise RuntimeError("The BRD analysis service could not complete this document. Please try again.") from last_error


def _draft_brief(sections, answers, fallback_summary):
    """Analyze scope, then develop each module into reviewable BRD specifications."""
    requirements = _section_items(sections)
    source = answers.get("source_requirement", "")
    # The structured requirements already contain the extracted source facts.
    # Including the entire raw document again made detailed requirements exceed
    # Groq's input limit before any BRD could be produced.
    context = {"reviewed_answers": answers.get("decisions", []),
               "requirements": requirements,
               "specialist_findings": answers.get("specialist_findings", {})}
    common = """You are Wooplix's dedicated BRD analyst. Produce a client-facing business
requirements document at the level of a detailed implementation BRD. Analyze the
actual project; never import facts from example clients. The original requirement
and reviewed answers are authoritative. A useful BRD expands a short requirement
into reviewable field definitions, business processes, rules, access, reports,
integrations and alerts where relevant. Distinguish explicit facts from a proposed
working design: propose reasonable details only within the requested capability.
Unknown mandatory flags, owners, thresholds, formulas, mappings, schedules and
recipients must be left blank. Never invent numeric targets, prices,
contractual commitments, software products or unrelated features. Preserve the
client's named products and scope boundaries. Write clear, specific business English.
Specialist findings are planning evidence. Use them only for proposed design
tables and risk/open-decision context; never convert a proposed finding into a
confirmed source requirement.
Specialist requirement IDs belong to a separate source register. Table references
must ONLY use the BR-xxx IDs supplied in this BRD call.
Return valid JSON only. Internal requirement IDs are for traceability; never put
them in titles, cells, or visible text. """
    overview_data = _groq_json(common + """Write an overview (100-180 words), a concise
business objective for every exact area, and cross-application tables only where
supported by the requirements. Schema: {"overview":"...",
"process_views":[{"area":"exact area", "description":"2-3 specific sentences",
"requirement_ids":["BR-001"]}], "specification_tables":[{"section":"Application Responsibilities",
"title":"Application responsibilities", "columns":["Application","Purpose","Key functions"],
"rows":[["...","...","..."]], "requirement_ids":["BR-001"]}]}.
Use Application Responsibilities, Objectives and Outcomes, Business Process and
Stakeholders as relevant. An outcome may be proposed but must not imply a measured
result. For process steps, make the sequence conditional when the client has not
confirmed it. Include one row per relevant application or business process stage.
Items requested only as cost categories are commercial questions, not delivery scope.
Avoid claims that implementation or support will be delivered when the client asks
only for their price. Do not describe the system as complete or guaranteed.
""", context, 2500)
    by_area = {}
    for item in requirements:
        by_area.setdefault(item["area"], set()).add(item["id"])
    views = []
    for view in overview_data.get("process_views", []):
        if not isinstance(view, dict):
            continue
        area, ids = str(view.get("area", "")).strip(), view.get("requirement_ids", [])
        if area in by_area and isinstance(ids, list) and ids and all(
                isinstance(x, str) and x in by_area[area] for x in ids):
            views.append({"area": area, "description": str(view.get("description", ""))[:900],
                          "requirement_ids": ids})
    tables = proposed_design_tables(overview_data.get("specification_tables", []), requirements, source)
    # Separate passes keep every product in view instead of letting early products
    # consume the model's output budget.
    substantive = [section for section in sections if section.get("requirements")
                   and str(section.get("product", "")).casefold() != "training"]
    for offset in range(0, len(substantive), 2):
        area_names = {row["product"] for row in substantive[offset:offset + 2]}
        area_requirements = [row for row in requirements if row["area"] in area_names]
        detail_data = _groq_json(common + """For ONLY the supplied business areas, create
substantive specification tables derived from their requirements. Do not just restate
the checklist. Give enough detail that stakeholders can review how each requested
capability would work. Use only relevant table types:
- Fields and Master Data: entity/field, suggested type or format, mandatory status,
  purpose/mapping. Suggest sensible field names when a form or record is requested;
  unknown type/mandatory/mapping must be blank.
- Workflows and Business Rules: trigger, condition, action, exception or decision.
- Roles and Access: role, record access, permitted action; leave unknown role names
  blank.
- Reports and Dashboards: KPI/widget, definition, data source, frequency/owner;
  unknown formulas and owners must be blank.
- Integrations: source, target, data exchanged, trigger or direction; only connect
  products explicitly requested to be integrated.
- Notifications: event/trigger, recipient, channel, content; leave unknowns
  blank. Do not invent a message template.
- Business Process: step, actor, action, output; proposed order is explicitly
  labelled as proposed, never as agreed.
Do not include sections unsupported by these areas. Include several meaningful
rows per applicable table and cover every substantive requirement in this group.
Prefer 3-8 rows per table. Use exact section names from: """ +
            ", ".join(SPECIFICATION_SECTIONS) + """.
Return {"specification_tables":[{"section":"Fields and Master Data",
"title":"Zoho CRM lead and account data", "columns":["Field","Type / format",
"Mandatory","Purpose"], "rows":[["...","...","","..."]],
"requirement_ids":["BR-001"]}]}.
Each table's requirement_ids must contain IDs supplied in this call. Proposed
details must be reviewable, specific, and restricted to the client's capability.
""", {"business_areas": list(area_names), "requirements": area_requirements,
        "reviewed_answers": answers.get("decisions", []),
        "specialist_findings": answers.get("specialist_findings", {})}, 2500)
        tables.extend(proposed_design_tables(detail_data.get("specification_tables", []),
                                             area_requirements, source))
        for area in area_names:
            area_ids = {r["id"] for r in area_requirements if r["area"] == area}
            area_tables = [table for table in tables
                           if set(table.get("requirement_ids", [])) & area_ids
                           and table["section"] != "Application Responsibilities"]
            if len(area_tables) >= 3 or (area.casefold() == "training" and area_tables):
                continue
            focused = [r for r in area_requirements if r["area"] == area]
            retry = _groq_json(common + """Develop this ONE missing business area into
specific, reviewable specification tables. Cover its requested capabilities.
Use relevant sections among Fields and Master Data, Workflows and Business Rules,
Roles and Access, Reports and Dashboards, Integrations, Notifications, Business
Process. Return {"specification_tables":[{"section":"...","title":"...",
"columns":["...","..."],"rows":[["...","..."]],
"requirement_ids":["BR-001"]}]}. Use only supplied IDs and source scope;
leave all unconfirmed values blank.""",
                                {"business_area": area, "requirements": focused,
                                "reviewed_answers": answers.get("decisions", [])}, 1800)
            tables.extend(proposed_design_tables(retry.get("specification_tables", []), focused, source))
            refined = [t for t in tables if set(t.get("requirement_ids", [])) & area_ids
                       and t["section"] != "Application Responsibilities"]
            if not refined:
                tables.append(_source_configuration_table(area, focused))
    if requirements and not tables:
        raise RuntimeError("The BRD analysis returned no usable detail. Please try again.")
    overview = str(overview_data.get("overview") or fallback_summary)[:1800]
    scope_text = " ".join(r["requirement"] for r in requirements).casefold()
    overview = " ".join(sentence.strip() for sentence in re.split(r"(?<=[.!?])\s+", overview)
                        if not any(term in sentence.casefold() and term not in scope_text
                                   for term in ("post-implementation support", "data migration")))
    return {"overview": overview or fallback_summary,
            "process_views": views, "specification_tables": tables}


def build_document(requirement, analysis, answers, source_name=""):
    """Make a source-grounded BRD with useful detail and no empty boilerplate."""
    all_questions = analysis.get("questions") or []
    questions = [q for q in all_questions if not _is_parallel_schedule_question(q.get("question", ""))]
    expected_answers = {q["id"] for q in questions}
    stale_schedule_answers = {q["id"] for q in all_questions if _is_parallel_schedule_question(q.get("question", ""))}
    if not expected_answers.issubset(set(answers)) or set(answers) - expected_answers - stale_schedule_answers:
        raise ValueError("Answer each question or choose Leave open.")
    for answer in answers.values():
        if not isinstance(answer, str) or not answer.strip() or len(answer) > 4000:
            raise ValueError("Enter each answer (up to 4000 characters).")

    sections = _combine_explicit_shared_areas(
        requirement, analysis.get("requirement_sections") or [])
    boundaries = []
    scoped_sections = []
    for section in sections:
        included = []
        for item in section.get("requirements", []):
            # Keep explicit future/excluded scope out of the committed checklist.
            if re.search(r"\b(?:not required (?:in|for|during) (?:the |this |current )?phase|out of scope|excluded from (?:this |the |current )?scope|future phase|future scope|future enhancement|optional feature)\b",
                         str(item), re.I):
                boundaries.append([str(section.get("product", "Business")), str(item)])
            else:
                included.append(item)
        if included:
            scoped_sections.append(dict(section, requirements=included))
    sections = scoped_sections
    requirements = _section_items(sections)
    decisions = [{"question": q["question"], "answer": answers[q["id"]]}
                 for q in questions]
    brief = _draft_brief(sections, {"source_requirement": requirement,
                                   "decisions": decisions,
                                   "specialist_findings": analysis.get("specialist_findings", {})},
                         str(analysis.get("summary") or "").strip())
    specifications = source_tables(requirement)
    # Preserve source definitions and add distinct proposed detail even when
    # both concern the same category (for example, two different forms).
    source_categories = {table["section"] for table in specifications}
    specifications += statement_tables(requirements, source_categories)
    existing = {(table["section"].casefold(), table["title"].casefold())
                for table in specifications}
    for table in brief.get("specification_tables", []):
        key = (table["section"].casefold(), table["title"].casefold())
        if key not in existing:
            specifications.append(table)
            existing.add(key)
    marketing = next((section for section in sections
                      if "marketing automation" in str(section.get("product", "")).casefold()), None)
    if marketing:
        requested = " ".join(str(item) for item in marketing.get("requirements", [])).casefold()
        drafted = " ".join(str(value) for table in specifications
                           if table.get("origin") == "proposed"
                           for row in table.get("rows", []) for value in row).casefold()
        missing = []
        if "journey" in requested and "journey" not in drafted:
            missing.append(["Marketing journey", "",
                            "Run the approved campaign sequence for the selected audience",
                            ""])
        if "follow-up automation" in requested and "follow-up" not in drafted:
            missing.append(["Follow-up automation", "",
                            "Create the agreed follow-up action after a campaign interaction",
                            ""])
        if missing:
            specifications.append({"section": "Workflows and Business Rules",
                                   "title": "Marketing journey and follow-up design",
                                   "columns": ["Capability", "Entry / trigger", "Proposed action", "Exit / exception"],
                                   "rows": missing, "origin": "proposed"})
    if boundaries:
        specifications.append({"section": "Scope Boundaries", "title": "Qualified scope",
                               "columns": ["Business area", "Source condition"],
                               "rows": boundaries, "origin": "source"})
    training_items = next((list(section.get("requirements", [])) for section in sections
                           if str(section.get("product", "")).casefold() == "training"), [])
    if training_items:
        plan_rows = []
        for item in training_items:
            name = str(item).strip()
            matched = [str(detail) for area in sections if area.get("product", "").casefold() != "training"
                       and (name.casefold() in area.get("product", "").casefold()
                            or (name.casefold() == "automation and reports" and
                                any(term in area.get("product", "").casefold()
                                    for term in ("analytics", "projects", "crm"))))
                       for detail in area.get("requirements", [])]
            focus = "; ".join(matched[:4]) if matched else "Practical use of the requested application"
            plan_rows.append([name, focus, "", ""])
        specifications.append({"section": "Training Plan", "title": "Practical training plan",
                               "columns": ["Application", "Practical focus", "Participants", "Duration"],
                               "rows": plan_rows, "origin": "proposed"})
    activities = []
    for section in sections:
        product = str(section.get("product") or "").strip()
        items = list(dict.fromkeys(str(x).strip() for x in section.get("requirements", []) if str(x).strip()))
        if items:
            activities.append({"area": product, "items": items})

    def matches(*terms):
        return [{"area": row["area"], "items": [item for item in row["items"]
                if any(term in item.casefold() for term in terms)]}
                for row in activities if any(any(term in item.casefold() for term in terms)
                                               for item in row["items"])]

    # Preserve stated objectives; only use this heading if it exists in the source.
    objectives = _source_list(requirement, "Main Objective", ["Cost Proposal", "Zoho CRM", "Training"])
    objectives += _source_list(requirement, "Business Objectives", ["Current Business", "Scope of Work"])
    purpose = _source_list(requirement, "Purpose", ["Scope", "Target Audience"])
    scope_items = _source_list(requirement, "Scope", ["Target Audience", "Current State Summary", "Business Context"])
    current_state = _source_list(requirement, "Current State Summary", ["Business Objectives", "Scope"])
    stakeholders = _source_list(requirement, "Target Audience", ["Current State Summary", "Business Objectives"])
    acceptance_criteria = (_source_list(requirement, "Go-Live Acceptance Criteria", ["Document Sign-Off"])
                           or _source_list(requirement, "Acceptance Criteria", ["Document Sign-Off"])
                           or _acceptance_table_rows(requirement))
    # Use the reviewed, source-filtered checklist as the authority for every product.
    areas = [{"name": row["area"], "requirements": row["items"]} for row in activities]
    open_decisions = []
    for question in questions:
        answer = answers[question["id"]].strip()
        if answer.casefold() not in {"leave open for discovery", "confirm during discovery", ""}:
            open_decisions.append({"question": question["question"], "answer": answer})

    requirements_by_area = [{"name": name} for name in dict.fromkeys(
        row["area"] for row in requirements) if name.casefold() != "training"]

    return {
        "title": "Business Requirements Document",
        "project_name": (Path(source_name).stem.replace("_", " ").strip()
                         if source_name and Path(source_name).stem.casefold() not in {"pasted requirement", "requirement"}
                         else "Project"),
        "prepared_for": "",
        "prepared_by": brand.COMPANY_NAME,
        "date": brand._ordinal_day(datetime.now(ZoneInfo("Asia/Kolkata"))),
        "status": "Draft for business review",
        "summary": brief["overview"],
        "purpose": purpose,
        "scope_items": scope_items,
        "process_views": brief["process_views"],
        "specification_tables": specifications,
        "objectives": objectives,
        "current_state": current_state,
        "stakeholders": stakeholders,
        "acceptance_criteria": acceptance_criteria,
        "areas": areas,
        "requirements": requirements,
        # These are tagged copies of source requirements, never newly authored
        # requirements. Rendering skips a category if it has no matching item.
        "business_rules": matches("rule", "duplicate", "same company", "approval", "mandatory", "validation", "billable / non-billable", "planned vs actual"),
        "data_requirements": matches("data", "record", "field", "duplicate", "event-wise", "batch-wise", "timesheet", "work notes", "master", "migration", "cleansing"),
        "integration_requirements": matches("integration", "integrat"),
        "reporting_requirements": matches("report", "dashboard", "analytics", "metric", "kpi"),
        "access_requirements": matches("role", "permission", "access", "security"),
        "requirements_by_area": requirements_by_area,
        "training_requirements": next((x["items"] for x in activities if x["area"].casefold() == "training"), []),
        "commercial_categories": analysis.get("commercial_categories") or [],
        "open_decisions": open_decisions,
        "source_name": Path(source_name).name if source_name else "",
        "activity_map": activities,
    }


def _e(value):
    return html.escape(_plain(value), quote=True)


def _plain(value):
    text = re.sub(r"\bBR-\d+\b", "", str(value or ""))
    text = re.sub(r"\bto be confirmed\b", "", text, flags=re.I)
    return re.sub(r"\s+", " ", text).strip()


def _specs_for(doc, categories):
    return [table for table in doc.get("specification_tables", [])
            if table["section"] in categories]


def _nonempty_areas(doc):
    return [area for area in doc["requirements_by_area"]
            if any(row["area"] == area["name"] for row in doc["requirements"])]


def build_html(doc):
    """Render a concise BRD; omit sections unsupported by the source request."""
    main = brand._logo_data_uri()
    badge = brand._badge_data_uri()
    main_img = f'<img src="{main}" style="position:absolute;left:0;top:0;height:38px" alt="Wooplix">' if main else ""
    badge_img = f'<img src="{badge}" style="position:absolute;right:0;top:0;height:26px" alt="Partner">' if badge else ""
    out = [f"""<!doctype html><html><head><meta charset="utf-8"><style>
@page {{ size:A4; margin:27mm 15mm 22mm; }}
body {{ font-family:'DejaVu Sans',sans-serif; font-size:9pt; line-height:1.45; color:#1e293b; }}
.header {{ position:fixed; top:-20mm; width:100%; z-index:1000; }} .header td:last-child {{text-align:right}}
.approval {{page-break-inside:avoid}}
.footer {{position:fixed;bottom:-8mm;width:100%;border-top:1px solid #cbd5e1;font-size:7pt;color:#64748b}}
.footer td {{padding-top:3mm}} h1 {{font-size:19pt;color:#002b49;margin:12mm 0 8mm}}
h2 {{page-break-after:avoid;background:#1a365d;color:#fff;font-size:11pt;padding:7px 10px;margin:18px 0 8px}}
h2.training {{page-break-before:always}}
h3 {{page-break-after:avoid;color:#008080;font-size:10pt;border-bottom:1px solid #008080;padding-bottom:3px;margin:13px 0 5px}}
p {{margin:4px 0 8px}} table {{width:100%;border-collapse:collapse;margin:6px 0 12px;font-size:8.5pt}}
thead {{display:table-header-group}} tr {{page-break-inside:avoid}} th {{background:#1a365d;color:white;text-align:left;font-size:8pt;padding:6px;border:1px solid #1a365d}}
td {{padding:5px 7px;vertical-align:top;border:1px solid #cbd5e1}} tbody tr:nth-child(even) td {{background:#f8fafc}}
.meta td:first-child {{background:#edf3f8;font-weight:bold;width:30%}} .muted {{color:#64748b;font-size:8pt}}
</style></head><body><div class="header">{main_img}{badge_img}</div>
<table class="footer"><tr><td><b>{_e(brand.COMPANY_NAME)}</b> · Confidential</td><td align="center">{_e(brand.COMPANY_EMAIL)}</td><td align="right">{_e(brand.COMPANY_WEBSITE)}</td></tr></table>
<h1>{_e(doc['title'])}</h1>
<table class="meta"><tbody>"""]
    for label, value in (("Project", doc["project_name"]), ("Status", doc["status"]),
                         ("Prepared for", doc.get("prepared_for", "")),
                         ("Prepared by", doc["prepared_by"]), ("Date", doc["date"])):
        if not value:
            continue
        out.append(f"<tr><td>{_e(label)}</td><td>{_e(value)}</td></tr>")
    out.append("</tbody></table>")

    chapter_number = 0

    def section(title):
        nonlocal chapter_number
        chapter_number += 1
        css_class = " class='training'" if title == "Training Requirements" else ""
        out.append(f"<h2{css_class}>{chapter_number}. {_e(title)}</h2>")

    def list_table(rows, first="Requirement"):
        out.append(f"<table><thead><tr><th style='width:8%'>No.</th><th>{_e(first)}</th></tr></thead><tbody>")
        for number, row in enumerate(rows, 1):
            out.append(f"<tr><td>{number}</td><td>{_e(row['requirement'])}</td></tr>")
        out.append("</tbody></table>")

    def specification_tables(category):
        tables = [table for table in doc.get("specification_tables", [])
                  if table["section"] == category]
        for table in tables:
            visible_columns = [i for i, name in enumerate(table["columns"])
                               if str(name).strip().casefold() not in
                               {"id", "brd id", "requirement id", "requirement ids", "traceability id"}]
            if not visible_columns:
                continue
            qualifier = " <span class='muted'>(Proposed for review)</span>" if table.get("origin") == "proposed" else ""
            out.append(f"<h3>{_e(table['title'])}{qualifier}</h3><table><thead><tr>")
            out.extend(f"<th>{_e(table['columns'][i])}</th>" for i in visible_columns)
            out.append("</tr></thead><tbody>")
            for row in table["rows"]:
                out.append("<tr>" + "".join(f"<td>{_e(row[i])}</td>" for i in visible_columns) + "</tr>")
            out.append("</tbody></table>")
        return bool(tables)

    section("Executive Summary")
    out.append(f"<p>{_e(doc['summary'])}</p>")
    if doc.get("purpose"):
        section("Purpose")
        out.append("<ul>" + "".join(f"<li>{_e(item)}</li>" for item in doc["purpose"]) + "</ul>")
    if doc["current_state"] or doc["objectives"] or doc["stakeholders"] or _specs_for(doc, {"Objectives and Outcomes"}):
        section("Business Context & Objectives")
    if doc["current_state"]:
        out.append("<h3>Current State</h3><ul>" + "".join(f"<li>{_e(item)}</li>" for item in doc["current_state"]) + "</ul>")
    if doc["objectives"]:
        out.append("<h3>Business Objectives</h3><ul>" + "".join(f"<li>{_e(item)}</li>" for item in doc["objectives"]) + "</ul>")
    if doc["stakeholders"]:
        out.append("<h3>Stakeholders</h3><ul>" + "".join(f"<li>{_e(item)}</li>" for item in doc["stakeholders"]) + "</ul>")
    specification_tables("Objectives and Outcomes")
    section("Project Scope")
    out.append("<ul>")
    scope_values = doc.get("scope_items") or [area["name"] for area in doc["requirements_by_area"]]
    for item in scope_values:
        out.append(f"<li>{_e(item)}</li>")
    out.append("</ul>")
    if _specs_for(doc, {"Scope Boundaries"}):
        out.append("<h3>Scope Conditions</h3>")
        specification_tables("Scope Boundaries")
    has_architecture = bool(_specs_for(doc, {"Application Responsibilities", "Business Process"}))
    if has_architecture:
        section("Solution Architecture")
    for category in ("Application Responsibilities",):
        if _specs_for(doc, {category}):
            out.append("<h3>Application Overview</h3>")
            specification_tables(category)
    if _specs_for(doc, {"Business Process"}):
        out.append("<h3>Business Process Flow</h3>")
        specification_tables("Business Process")
    if _specs_for(doc, {"Fields and Master Data"}):
        section("Master Data Requirements")
        specification_tables("Fields and Master Data")

    section("Functional Requirements by Module")
    rows_by_area = {}
    for row in doc["requirements"]:
        rows_by_area.setdefault(row["area"], []).append(row)
    for area in _nonempty_areas(doc):
        name = area["name"]
        out.append(f"<h3>{_e(name)}</h3>")
        process_view = next((view for view in doc["process_views"] if view["area"] == name), None)
        if process_view:
            out.append(f"<h4>Business Objective</h4><p>{_e(process_view['description'])}</p>")
        rows = rows_by_area.get(name, [])
        if rows:
            out.append("<table><thead><tr><th style='width:8%'>No.</th><th>Business Requirement</th></tr></thead><tbody>")
            for number, row in enumerate(rows, 1):
                out.append(f"<tr><td>{number}</td><td>{_e(row['requirement'])}</td></tr>")
            out.append("</tbody></table>")
    for title, categories in (
        ("Approval Workflow & Authorization Matrix", ("Workflows and Business Rules", "Roles and Access")),
        ("Dashboard Specifications", ("Reports and Dashboards",)),
        ("Integrations", ("Integrations",)),
        ("Alerting & Notification System", ("Notifications",)),
        ("Additional Business Details", ("Business Details",)),
    ):
        tables = _specs_for(doc, categories)
        if tables:
            section(title)
            for category in categories:
                specification_tables(category)
    if _specs_for(doc, {"Stakeholders"}):
        section("Stakeholder Roles")
        specification_tables("Stakeholders")
    if doc["training_requirements"]:
        section("Training Requirements")
        if _specs_for(doc, {"Training Plan"}):
            specification_tables("Training Plan")
        else:
            list_table([row for row in doc["requirements"] if row["area"].casefold() == "training"], "Training requirement")
    if doc["open_decisions"]:
        section("Questions and Decisions")
        out.append("<table><thead><tr><th>Question</th><th>Answer or status</th></tr></thead><tbody>")
        for row in doc["open_decisions"]:
            out.append(f"<tr><td>{_e(row['question'])}</td><td>{_e(row['answer'])}</td></tr>")
        out.append("</tbody></table>")
    if any(t["section"] == "Acceptance Criteria" for t in doc.get("specification_tables", [])):
        section("Acceptance Criteria")
        specification_tables("Acceptance Criteria")
    elif doc["acceptance_criteria"]:
        section("Acceptance Criteria")
        out.append("<ul>" + "".join(f"<li>{_e(item)}</li>" for item in doc["acceptance_criteria"]) + "</ul>")
    from multi_agent_tools import planning_tables
    for title, columns, rows in planning_tables(doc):
        section(title)
        out.append('<table><thead><tr>' + ''.join(f'<th>{_e(column)}</th>' for column in columns) +
                   '</tr></thead><tbody>')
        for values in rows:
            out.append('<tr>' + ''.join(f'<td>{_e(value)}</td>' for value in values) + '</tr>')
        out.append('</tbody></table>')
    out.append("</body></html>")
    return "".join(out)


def build_docx(doc, output):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    report = Document()
    section = report.sections[0]
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.left_margin = section.right_margin = Inches(0.7)
    section.top_margin = section.bottom_margin = Inches(0.75)
    normal = report.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
    header = section.header
    logo_table = header.add_table(rows=1, cols=2, width=Inches(7.1))
    if Path(brand.LOGO_MAIN_PATH).exists():
        logo_table.cell(0, 0).paragraphs[0].add_run().add_picture(brand.LOGO_MAIN_PATH, width=Inches(2.0))
    if Path(brand.LOGO_BADGE_PATH).exists():
        p_badge = logo_table.cell(0, 1).paragraphs[0]
        p_badge.alignment = 2
        p_badge.add_run().add_picture(brand.LOGO_BADGE_PATH, width=Inches(1.65))
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer.add_run(f"{brand.COMPANY_NAME}  ·  Confidential  ·  {brand.COMPANY_WEBSITE}").font.size = Pt(8)
    p = report.add_paragraph()
    p.alignment = 1
    r = p.add_run(_plain(doc["title"]))
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(0, 43, 73)
    table = report.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in (("Project", doc["project_name"]), ("Status", doc["status"]),
                         ("Prepared for", doc.get("prepared_for", "")),
                         ("Prepared by", doc["prepared_by"]), ("Date", doc["date"])):
        if not value:
            continue
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, _plain(value)
        cells[0]._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="edf3f8"/>'))

    chapter_number = 0

    def heading(value):
        nonlocal chapter_number
        chapter_number += 1
        h = report.add_heading(f"{chapter_number}. {value}", level=1)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0x1a, 0x36, 0x5d)
        return h

    def bullets(values):
        for value in values:
            report.add_paragraph(str(value), style="List Bullet")

    def specification_tables(category):
        tables = [table for table in doc.get("specification_tables", [])
                  if table["section"] == category]
        for spec in tables:
            title = _plain(spec["title"])
            if spec.get("origin") == "proposed":
                title += " (Proposed for review)"
            report.add_heading(title, level=2)
            visible_columns = [i for i, name in enumerate(spec["columns"])
                               if str(name).strip().casefold() not in
                               {"id", "brd id", "requirement id", "requirement ids", "traceability id"}]
            if not visible_columns:
                continue
            table = report.add_table(rows=1, cols=len(visible_columns))
            table.style = "Table Grid"
            for cell, index in zip(table.rows[0].cells, visible_columns):
                cell.text = _plain(spec["columns"][index])
            for values in spec["rows"]:
                for cell, index in zip(table.add_row().cells, visible_columns):
                    cell.text = _plain(values[index])
        return bool(tables)

    heading("Executive Summary")
    report.add_paragraph(_plain(doc["summary"]))
    if doc.get("purpose"):
        heading("Purpose")
        bullets([_plain(item) for item in doc["purpose"]])
    if doc["current_state"] or doc["objectives"] or doc["stakeholders"] or _specs_for(doc, {"Objectives and Outcomes"}):
        heading("Business Context & Objectives")
    if doc["current_state"]:
        report.add_heading("Current State", level=2)
        bullets([_plain(item) for item in doc["current_state"]])
    if doc["objectives"]:
        report.add_heading("Business Objectives", level=2)
        bullets([_plain(item) for item in doc["objectives"]])
    if doc["stakeholders"]:
        report.add_heading("Stakeholders", level=2)
        bullets([_plain(item) for item in doc["stakeholders"]])
    specification_tables("Objectives and Outcomes")
    heading("Project Scope")
    scope_values = doc.get("scope_items") or [area["name"] for area in doc["requirements_by_area"]]
    bullets([_plain(item) for item in scope_values])
    if _specs_for(doc, {"Scope Boundaries"}):
        report.add_heading("Scope Conditions", level=2)
        specification_tables("Scope Boundaries")
    has_architecture = bool(_specs_for(doc, {"Application Responsibilities", "Business Process"}))
    if has_architecture:
        heading("Solution Architecture")
    if _specs_for(doc, {"Application Responsibilities"}):
        report.add_heading("Application Overview", level=2)
        specification_tables("Application Responsibilities")
    if _specs_for(doc, {"Business Process"}):
        report.add_heading("Business Process Flow", level=2)
        specification_tables("Business Process")
    if _specs_for(doc, {"Fields and Master Data"}):
        heading("Master Data Requirements")
        specification_tables("Fields and Master Data")

    heading("Functional Requirements by Module")
    rows_by_area = {}
    for row in doc["requirements"]:
        rows_by_area.setdefault(row["area"], []).append(row)
    for area in _nonempty_areas(doc):
        name = area["name"]
        report.add_heading(_plain(name), level=2)
        process_view = next((view for view in doc["process_views"] if view["area"] == name), None)
        if process_view:
            report.add_heading("Business Objective", level=3)
            report.add_paragraph(_plain(process_view["description"]))
        rows = rows_by_area.get(name, [])
        if rows:
            table = report.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "No."
            table.rows[0].cells[1].text = "Business Requirement"
            for number, row in enumerate(rows, 1):
                cells = table.add_row().cells
                cells[0].text, cells[1].text = str(number), _plain(row["requirement"])

    for title, categories in (
        ("Approval Workflow & Authorization Matrix", ("Workflows and Business Rules", "Roles and Access")),
        ("Dashboard Specifications", ("Reports and Dashboards",)),
        ("Integrations", ("Integrations",)),
        ("Alerting & Notification System", ("Notifications",)),
        ("Additional Business Details", ("Business Details",)),
    ):
        if _specs_for(doc, categories):
            heading(title)
            for category in categories:
                specification_tables(category)
    if _specs_for(doc, {"Stakeholders"}):
        heading("Stakeholder Roles")
        specification_tables("Stakeholders")

    if doc["training_requirements"]:
        training_heading = heading("Training Requirements")
        training_heading.paragraph_format.page_break_before = True
        if _specs_for(doc, {"Training Plan"}):
            specification_tables("Training Plan")
        else:
            rows = [row for row in doc["requirements"] if row["area"].casefold() == "training"]
            table = report.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text, table.rows[0].cells[1].text = "No.", "Training Requirement"
            for number, row in enumerate(rows, 1):
                cells = table.add_row().cells
                cells[0].text, cells[1].text = str(number), _plain(row["requirement"])
    if doc["open_decisions"]:
        heading("Questions and Decisions")
        table = report.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Question", "Answer or status"
        for row in doc["open_decisions"]:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = _plain(row["question"]), _plain(row["answer"])
    if any(t["section"] == "Acceptance Criteria" for t in doc.get("specification_tables", [])):
        heading("Acceptance Criteria")
        specification_tables("Acceptance Criteria")
    elif doc["acceptance_criteria"]:
        heading("Acceptance Criteria")
        bullets([_plain(item) for item in doc["acceptance_criteria"]])
    from multi_agent_tools import planning_tables
    for title, columns, rows in planning_tables(doc):
        heading(title)
        table = report.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        for cell, column in zip(table.rows[0].cells, columns):
            cell.text = column
        for values in rows:
            for cell, value in zip(table.add_row().cells, values):
                cell.text = _plain(value)
    # Match the Wooplix table palette and repeat column labels across pages.
    for index, table in enumerate(report.tables):
        if index:
            first = table.rows[0]
            first._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            for cell in first.cells:
                cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="1a365d"/>'))
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
        for row in table.rows:
            row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    report.save(output)
    return output


def build_pdf(doc, output):
    php = shutil.which("php")
    script = Path(brand.PHP_SCRIPT)
    if not php or not (Path(brand.HERE) / "vendor" / "autoload.php").exists():
        raise RuntimeError("The PDF renderer is unavailable.")
    fd, html_path = tempfile.mkstemp(suffix=".html", prefix="wooplix_brd_", dir=str(Path(output).parent))
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            stream.write(build_html(doc))
        result = subprocess.run([php, str(script), html_path, str(output)],
                                capture_output=True, text=True, timeout=240)
        if result.returncode or not Path(output).exists():
            raise RuntimeError("Could not render the Business Requirements Document PDF.")
    finally:
        try:
            os.unlink(html_path)
        except OSError:
            pass
    return output
